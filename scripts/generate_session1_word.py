from pathlib import Path
from docx import Document

preferred_out_path = Path("c:/Users/cmari/Desktop/IEP/CriptografíaPostcuantica/05_entregables/02_word/sesion_1_fundamentos_criptografia.docx")
backup_out_path = preferred_out_path.with_name("sesion_1_fundamentos_criptografia_refinado.docx")
preferred_out_path.parent.mkdir(parents=True, exist_ok=True)

doc = Document()

doc.add_heading("INTRODUCCIÓN A LA CRIPTOGRAFÍA: FUNDAMENTOS, MODELO DE AMENAZA Y CRIPTOGRAFÍA MODERNA", 0)
doc.add_paragraph("TEMA 1")
doc.add_heading("Fundamentos de la criptografía", 1)

def add_section(title, paragraphs):
    doc.add_heading(title, level=2)
    for p in paragraphs:
        doc.add_paragraph(p)


def add_subsection(title, paragraphs=None, bullets=None):
    doc.add_heading(title, level=3)
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if bullets:
        for item in bullets:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(item)

add_section(
    "1. Introducción",
    [
        "La criptografía es una disciplina de la seguridad de la información que tiene como finalidad proteger la confidencialidad, la integridad, la autenticación y la autenticidad de la información cuando se transmite o se almacena en entornos potencialmente inseguros.",
        "No consiste únicamente en ocultar mensajes. La criptografía moderna estudia técnicas matemáticas y protocolos para proporcionar propiedades de seguridad incluso cuando los sistemas se comunican a través de canales controlados por un atacante.",
        "Desde la antigüedad hasta la actualidad, la necesidad de ocultar información y verificar la identidad de los interlocutores ha sido una preocupación constante. En la práctica digital, la criptografía aparece en la navegación web, en la banca online, en la firma de documentos, en la autenticación de usuarios y en la protección de datos en la nube.",
        "La seguridad de la información no depende solo del ocultamiento del contenido, sino también del control del acceso, la autenticación de las entidades, la verificación de la integridad y la posibilidad de acreditar la procedencia de un mensaje.",
        "Fecha de referencia: XX-YY-ZZZZ"
    ],
)

add_section(
    "2. Modelo de amenaza: Alice, Bob y Eve",
    [
        "Antes de estudiar algoritmos, es necesario entender el escenario de seguridad. En un modelo clásico, Alice quiere enviar un mensaje a Bob a través de un canal que puede ser inseguro, mientras que Eve representa a un atacante que puede escuchar, modificar, inyectar o eliminar mensajes.",
        "Existen dos tipos principales de atacante. Un atacante pasivo escucha el canal sin alterar los mensajes. Un atacante activo puede interceptar el tráfico, modificar contenido, repetir mensajes previamente capturados o suplantar identidades.",
        "Esto da lugar al ataque de intermediario o MITM, en el que Alice cree que habla con Bob y Bob cree que habla con Alice, pero Eve controla el canal de comunicación. Este tipo de ataque explica por qué no basta con cifrar el mensaje: también es necesario autenticar a las entidades y proteger la integridad del flujo de datos.",
        "La seguridad computacional es diferente de la seguridad absoluta. En la práctica, se acepta que un sistema es seguro si un atacante con recursos razonables no puede romperlo de forma práctica. Es decir, la seguridad se fundamenta en problemas computacionalmente difíciles y en la correcta gestión de los mecanismos criptográficos."
    ],
)

add_section(
    "3. Objetivos de la seguridad de la información",
    [
        "Los objetivos de la seguridad suelen describirse de la siguiente forma:",
        "Confidencialidad: ¿quién puede leer el mensaje? La información solo puede ser entendida por quienes están autorizados para acceder a ella.",
        "Integridad: ¿ha sido modificado? El contenido no debe cambiar ni durante la transmisión ni durante el almacenamiento.",
        "Autenticación de entidad: ¿con quién estoy hablando? Debe verificarse la identidad de la entidad que aparece en la comunicación.",
        "Autenticación de mensaje: ¿quién generó este mensaje? Debe poder determinarse si el mensaje procede de la entidad que afirma ser su origen.",
        "No repudio: ¿puede una parte negar la autoría o recepción de un mensaje? Esto depende no solo de la criptografía, sino también del protocolo, la gestión de claves y el contexto legal.",
        "Además, la disponibilidad es un objetivo general de la seguridad, aunque no sea una propiedad que proporcione directamente la criptografía. El control de acceso tampoco debe confundirse con la criptografía, aunque ambas son componentes esenciales de la seguridad de un sistema."
    ],
)

add_section(
    "4. ¿Qué es la criptografía?",
    [
        "La criptografía es el conjunto de técnicas matemáticas y protocolos diseñados para proporcionar propiedades de seguridad como confidencialidad, integridad, autenticación y no repudio, incluso en entornos adversariales.",
        "Su campo de estudio incluye el cifrado, las funciones hash, los códigos de autenticación de mensajes, las firmas digitales, los protocolos de intercambio de claves y la gestión segura de las claves.",
        "Desde un punto de vista operativo, la criptografía transforma un mensaje original en una representación protegida, presenta la información a través de un canal inseguro y recupera el contenido original solo con la clave o el procedimiento apropiado.",
        "En la práctica, la seguridad de un sistema no depende solo del algoritmo elegido. También influyen la gestión de claves, la longitud de la clave, la aleatoriedad, la implementación, los protocolos y la correcta elección de parámetros como nonces o vectores de inicialización."
    ],
)

add_section(
    "5. Primitivas criptográficas",
    [
        "Las primitivas criptográficas son los bloques básicos sobre los que se construyen sistemas más complejos. Entenderlas es esencial, porque todas las soluciones reales de seguridad (TLS, certificados, firma digital, almacenamiento seguro, autenticación, cifrado de datos) se apoyan en estas ideas.",
        "Una forma clara de organizarlas es distinguir entre cinco tipos principales: cifrado, funciones hash, MAC, firmas digitales e intercambio de claves."
    ],
)

add_subsection(
    "5.1 Cifrado",
    paragraphs=[
        "El cifrado transforma un mensaje M con una clave K en un texto cifrado C. La idea principal es que solo alguien que conozca la clave correcta pueda recuperar el contenido original.",
        "El objetivo principal del cifrado es la confidencialidad: impedir que un atacante comprenda el contenido del mensaje aunque lo intercepte."
    ],
    bullets=[
        "Se usa para ocultar el contenido del mensaje desde el punto de vista del atacante.",
        "En la criptografía simétrica, se usa la misma clave para cifrar y descifrar.",
        "En la asimétrica, se usa la clave pública del destinatario para cifrar y su clave privada para descifrar."
    ],
)

add_subsection(
    "5.2 Funciones hash",
    paragraphs=[
        "Una función hash transforma cualquier entrada en una salida de longitud fija, conocida como digest o valor hash. El proceso es determinista y no reversible en la práctica."
    ],
    bullets=[
        "Se usan para verificar integridad de archivos, mensajes y contraseñas.",
        "Si cambia un solo bit del contenido original, el hash cambia de forma muy distinta.",
        "No permiten recuperar el mensaje original a partir del valor hash."
    ],
)

add_subsection(
    "5.3 MAC (Message Authentication Code)",
    paragraphs=[
        "Un MAC permite comprobar la integridad y la autenticidad de un mensaje entre dos entidades que comparten una clave secreta. Es decir, confirma que el mensaje no ha sido modificado y que procede de quien posee la clave compartida."
    ],
    bullets=[
        "Se utiliza cuando existe una clave compartida entre emisor y receptor.",
        "Protege frente a modificaciones no autorizadas del mensaje.",
        "No es lo mismo que una firma digital, porque la verificación exige compartir la clave secreta."
    ],
)

add_subsection(
    "5.4 Firmas digitales",
    paragraphs=[
        "Una firma digital permite demostrar que un mensaje ha sido firmado por la persona que afirma ser su autor y que el contenido no ha sido alterado desde la firma. Se basa en un par de claves: pública y privada."
    ],
    bullets=[
        "La clave privada firma el mensaje o su resumen.",
        "La clave pública valida la firma.",
        "Proporcionan autenticidad, integridad y, en muchos contextos, no repudio."
    ],
)

add_subsection(
    "5.5 Intercambio o acuerdo de claves",
    paragraphs=[
        "El intercambio o acuerdo de claves permite que dos partes establezcan un secreto compartido sin enviar ese secreto directamente por el canal de comunicación."
    ],
    bullets=[
        "Diffie-Hellman es el ejemplo clásico de protocolo de acuerdo de claves.",
        "Ambas partes calculan un valor común sin revelarlo en claro.",
        "En la práctica, se combina con autenticación para evitar ataques de intermediario."
    ],
)

add_section(
    "Resumen práctico",
    [
        "El cifrado protege la confidencialidad del contenido; los hashes ayudan a detectar cambios; los MAC autentican mensajes compartiendo una clave; las firmas digitales permiten demostrar origen e integridad; y el acuerdo de claves permite crear secretos compartidos sin enviarlos directamente."
    ],
)

add_section(
    "6. Hash frente a cifrado",
    [
        "Es importante distinguir claramente entre hash y cifrado. El cifrado es reversible: con la clave adecuada se puede recuperar el mensaje original. La función hash no es reversible en la práctica: a partir del hash no se debe poder recuperar el contenido original.",
        "Ejemplo: 'Hola' -> SHA-256 -> digest. Esto no significa que el digest sea una versión cifrada de 'Hola'; simplemente es una huella digital del contenido. Si se cambia una sola letra, el hash cambia por completo.",
        "Mecanismo: cifrado. ¿Se puede recuperar el mensaje original? Sí. ¿Usa clave? Sí. Uso principal: confidencialidad.",
        "Mecanismo: hash. ¿Se puede recuperar el mensaje original? No. ¿Usa clave? No. Uso principal: integridad y construcción de primitivas.",
        "Mecanismo: MAC. ¿Se puede recuperar el mensaje original? No. ¿Usa clave? Sí. Uso principal: integridad y autenticidad.",
        "Mecanismo: firma digital. ¿Se puede recuperar el mensaje original? No. ¿Usa clave? Sí, par público/privado. Uso principal: autenticidad e integridad."
    ],
)

add_section(
    "7. Criptografía simétrica",
    [
        "La criptografía simétrica utiliza una sola clave compartida entre el emisor y el receptor. Es rápida y eficiente para cifrar grandes volúmenes de datos, por lo que se usa mucho para proteger el contenido de mensajes o archivos.",
        "AES es uno de los algoritmos simétricos más relevantes. Se trata de un cifrado por bloques, no de un sistema completo por sí solo para cifrar cualquier flujo de datos de forma directa. En la práctica, se combina con modos de operación como ECB, CBC, CTR y GCM.",
        "ECB es un modo muy simple, pero problemático: bloques iguales producen bloques cifrados iguales. Eso puede revelar patrones en el contenido cifrado. CBC y CTR mejoran esa situación, y GCM combina cifrado con autenticación, proporcionando confidencialidad, integridad y autenticidad del texto cifrado.",
        "El cifrado autenticado o AEAD (Authenticated Encryption with Associated Data) es una de las grandes ideas de la criptografía moderna. Ejemplos claros incluyen AES-GCM y ChaCha20-Poly1305. En estos sistemas se protege la confidencialidad y la integridad del mensaje de forma conjunta.",
        "Los nonces, IVs y la aleatoriedad son elementos críticos. El mismo nonce no debe reutilizarse con la misma clave en esquemas como GCM. La buena gestión de claves, nonces y valores aleatorios es tan importante como el algoritmo escogido."
    ],
)

add_section(
    "8. Criptografía asimétrica",
    [
        "La criptografía asimétrica se basa en un par de claves. La clave pública se puede distribuir libremente, mientras que la clave privada debe mantenerse en secreto. Es la base de la autenticación y del intercambio seguro de claves.",
        "El cifrado con clave pública permite que el destinatario reciba un mensaje que solo él pueda descifrar con su clave privada. La firma digital permite que el emisor firme un mensaje con su clave privada y que cualquier usuario valide la firma con la clave pública asociada.",
        "Los principales métodos asimétricos incluyen RSA, ECC y ElGamal. Si bien RSA es muy conocido y ampliamente usado, ECC ofrece un nivel de seguridad equivalente con claves mucho más pequeñas, por lo que es especialmente relevante en dispositivos móviles, sistemas embebidos y protocolos modernos.",
        "Es fundamental distinguir entre tres usos diferentes: cifrado de clave pública, firma digital y acuerdo de claves. No todos se resuelven con la misma operación ni con la misma lógica interna."
    ],
)

add_section(
    "9. Firmas digitales y la explicación correcta",
    [
        "La firma digital no es simplemente 'cifrar el hash con la clave privada' de forma mecánica. La mejor forma de entenderla es como una operación criptográfica que se calcula sobre un mensaje, normalmente sobre un hash del mismo, usando la clave privada. La verificación se hace con la clave pública correspondiente.",
        "Este enfoque permite verificar dos cosas: que el emisor era quien decía ser y que el mensaje no ha sido alterado desde la firma. Modelos modernos incluyen RSA-PSS, Ed25519 y ECDSA. La clave es que la firma digital aporta autenticidad e integridad, no solo ocultación.",
        "En otras palabras, la firma digital es un mecanismo de prueba criptográfica de que el mensaje procede del emisor correcto y de que el contenido es el mismo que el que se firmó originalmente."
    ],
)

add_section(
    "10. Diffie-Hellman y el acuerdo de claves",
    [
        "Diffie-Hellman es un protocolo clave para establecer un secreto compartido sin enviar ese secreto directamente por el canal. Alice y Bob acuerdan públicamente un valor p y un generador g.",
        "Alice elige un valor secreto a y calcula A = g^a mod p. Bob elige b y calcula B = g^b mod p. Ambos intercambian A y B. Luego Alice calcula K = B^a mod p y Bob calcula K = A^b mod p. Ambos obtienen el mismo valor K = g^(ab) mod p.",
        "Este mecanismo es muy útil, pero por sí solo no autentica a Alice ni a Bob. Sin autenticación, un atacante puede montarse en el medio y realizar un ataque de intermediario. Por eso, en la práctica, Diffie-Hellman se combina con autenticación, certificados o mecanismos similares."
    ],
)

add_section(
    "11. Fundamentos matemáticos para la criptografía",
    [
        "La criptografía moderna se apoya en herramientas matemáticas como la aritmética modular, la teoría de números, la exponenciación modular, el inverso modular, el máximo común divisor, la factorización y el logaritmo discreto.",
        "La congruencia es la base de estos sistemas: a ≡ b (mod n) significa que a y b tienen el mismo resto al dividir entre n. Ejemplo: 23 ≡ 2 (mod 7), porque 23 = 3 × 7 + 2.",
        "La exponenciación modular es esencial en RSA y Diffie-Hellman. Por ejemplo, 3^4 mod 5 = 1, porque 81 ≡ 1 (mod 5). El algoritmo de Euclides y el cálculo del inverso modular preparan al estudiante para entender la generación de claves y la operación de cifrado y descifrado en sistemas asimétricos.",
        "Los números primos y la función φ de Euler son también esenciales. En RSA, se toma n = p × q y se usa φ(n) = (p - 1)(q - 1) para construir las claves. El detalle técnico puede ser guiado por la noción de que la factorización de un número grande es un problema difícil para un ordenador clásico."
    ],
)

add_section(
    "12. Ejemplo didáctico de RSA",
    [
        "Para entender la idea, usamos un ejemplo pequeño. Supongamos p = 5 y q = 11. Entonces n = 5 × 11 = 55 y φ(n) = (5 - 1)(11 - 1) = 40.",
        "Elegimos e = 3. Buscamos d tal que e × d ≡ 1 (mod 40). Por ejemplo, 3 × 27 = 81 ≡ 1 (mod 40). Con esto, la clave pública es (55, 3) y la clave privada es 27.",
        "Para cifrar un mensaje m, se usa c = m^e mod n. Para descifrar, se usa m = c^d mod n. Este ejemplo es solo educativo: en sistemas reales no se implementa RSA a mano ni sin padding seguro. Para uso real se usan esquemas como RSA-OAEP para cifrado y RSA-PSS para firmas."
    ],
)

add_section(
    "13. Ejemplos prácticos",
    [
        "Ejemplo 1: Si Alice cifra un mensaje con la clave pública de Bob, Bob puede descifrarlo con su clave privada. Esto resuelve la confidencialidad del contenido.",
        "Ejemplo 2: Si Alice firma un documento con su clave privada, Bob puede verificar la firma con la clave pública de Alice. Esto resuelve la autenticidad y la integridad del documento.",
        "Ejemplo 3: Si Alice y Bob usan Diffie-Hellman sin autenticación, Eve puede actuar como intermediario y negociar un secreto distinto con cada uno. Esto muestra la importancia de autenticar las claves o identidades.",
        "Estos ejemplos ayudan a comprender que la criptografía no es solo una herramienta de ocultación, sino un conjunto de mecanismos para construir comunicaciones seguras y verificables."
    ],
)

add_section(
    "14. Ejercicios propuestos",
    [
        "Ejercicio 1: Explica la diferencia entre criptografía simétrica y asimétrica.",
        "Solución orientativa: la simétrica usa una misma clave para cifrar y descifrar, mientras que la asimétrica usa un par de claves públicas y privadas.",
        "Ejercicio 2: ¿Qué implica la expresión 23 ≡ 2 (mod 7)?",
        "Solución orientativa: ambos números tienen el mismo resto al dividir entre 7.",
        "Ejercicio 3: ¿Qué diferencia hay entre un hash y un cifrado?",
        "Solución orientativa: el cifrado es reversible y usa clave; el hash es unidireccional y sirve para integridad.",
        "Ejercicio 4: ¿Qué problema resuelve Diffie-Hellman?",
        "Solución orientativa: permite acordar una clave compartida sin enviarla directamente por el canal.",
        "Ejercicio 5: ¿Por qué es importante distinguir entre MAC y firma digital?",
        "Solución orientativa: un MAC autentica mensajes para quienes comparten la clave; la firma digital usa un par público/privado y permite verificación por terceros."
    ],
)

add_section(
    "15. Python como apoyo didáctico",
    [
        "Python permite introducir conceptos clave con poca complejidad. Un ejemplo sencillo de función hash y aritmética modular es el siguiente:",
        "import hashlib\nmensaje = b'Hola, mundo seguro'\nprint(hashlib.sha256(mensaje).hexdigest())\nprint(23 % 7)\nprint(pow(3, 4, 5))",
        "Estos ejercicios ilustran que es posible comprobar integridad, calcular restos y trabajar con potencias modulares de forma accesible para principiantes.",
        "La clave didáctica es conectar la matemática con ejemplos de seguridad real sin perder claridad conceptual."
    ],
)

add_section(
    "16. Importancia de la seguridad digital",
    [
        "La criptografía es una pieza central de la seguridad digital. Sostiene servicios como la navegación web, la banca online, la gestión de credenciales, la firma de documentos, la seguridad de correos electrónicos y la autenticación en redes empresariales.",
        "Sin una base criptográfica sólida, las organizaciones y los individuos quedan expuestos a robo de datos, manipulación, suplantación de identidad y fraudes digitales.",
        "La seguridad real de un sistema no se logra con un único algoritmo sino mediante la combinación de objetivos de seguridad, protocolos correctos, gestión adecuada de claves, autenticación y buenas prácticas de implementación."
    ],
)

add_section(
    "17. Errores comunes y buenas prácticas",
    [
        "Un algoritmo robusto puede resultar inseguro si se emplea de forma incorrecta. Errores frecuentes incluyen la reutilización de claves o nonces, el uso de algoritmos obsoletos, la mala gestión de contraseñas y la falta de autenticación.",
        "Las buenas prácticas incluyen la generación de claves con fuentes aleatorias criptográficamente seguras, la validación de certificados, el uso de KDF para contraseñas, el uso de AEAD en protocolos modernos y la prohibición de diseñar algoritmos propios.",
        "Como regla general, se recomienda utilizar librerías criptográficas revisadas y protocolos ampliamente evaluados por la comunidad, en lugar de desarrollar soluciones ad hoc."
    ],
)

add_section(
    "18. Conclusiones",
    [
        "La criptografía moderna no es solo una herramienta para ocultar mensajes, sino un conjunto de mecanismos para proteger la seguridad de la información ante ataques reales.",
        "Los conceptos más importantes de esta sesión son la confidencialidad, la integridad, la autenticación, la gestión segura de claves y el papel del atacante en el modelo de amenaza.",
        "Con una base sólida de matemáticas, primitivas y protocolos, el alumno puede comprender mejor la lógica que sustenta los sistemas de seguridad actuales, desde TLS hasta las firmas digitales y la protección de datos en entornos digitales."
    ],
)

add_section(
    "19. Bibliografía esencial",
    [
        "Stallings, W. Cryptography and Network Security: Principles and Practice. Pearson.",
        "Schneier, B. Applied Cryptography. Wiley.",
        "Katz, J. y Lindell, Y. Introduction to Modern Cryptography. Chapman & Hall/CRC.",
        "Menezes, A., van Oorschot, P. y Vanstone, S. Handbook of Applied Cryptography. CRC Press.",
        "National Institute of Standards and Technology (NIST). Digital Identity Guidelines and related cryptographic standards documentation."
    ],
)

add_section(
    "20. Bibliografía recomendada",
    [
        "Buchmann, J. A. Introduction to Cryptography. Springer.",
        "Silverman, J. H. A Friendly Introduction to Number Theory. Pearson.",
        "Ferguson, N., Schneier, B. y Kohno, T. Cryptography Engineering. Wiley.",
        "Documentación oficial de TLS, X.509, PKI, AES-GCM, Diffie-Hellman y protocolos de firma digital."
    ],
)

add_section(
    "21. Preguntas de reflexión",
    [
        "¿Por qué no basta con cifrar un mensaje si no se autentica al interlocutor?",
        "¿Qué diferencia hay entre confidencialidad e integridad?",
        "¿Qué problema resuelve Diffie-Hellman y qué limita en su versión básica?",
        "¿Por qué la gestión de claves y nonces es crítica en protocolos modernos?",
        "¿Qué papel juega la firma digital en la seguridad de documentos y transacciones?",
        "¿Qué ocurre si se reutiliza una clave o un nonce en un esquema como AES-GCM?"
    ],
)

out_path = preferred_out_path
try:
    doc.save(out_path)
    print(f"Word generated: {out_path}")
except PermissionError:
    out_path = backup_out_path
    doc.save(out_path)
    print(f"Word generated: {out_path}")
    print("The main output file was locked by another application. A refined copy was generated instead.")
