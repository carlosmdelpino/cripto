from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

preferred_out_path = Path("c:/Users/cmari/Desktop/IEP/CriptografíaPostcuantica/05_entregables/02_word/sesion_3_computacion_cuantica.docx")
backup_out_path = preferred_out_path.with_name("sesion_3_computacion_cuantica_v2.docx")
preferred_out_path.parent.mkdir(parents=True, exist_ok=True)

doc = Document()

doc.add_heading("INTRODUCCIÓN A LA COMPUTACIÓN CUÁNTICA: QUBITS, ALGORITMOS Y EL IMPACTO SOBRE LA CRIPTOGRAFÍA", 0)
doc.add_paragraph("TEMA 3 — Versión reducida (sesión de 30 minutos)")
doc.add_heading("Computación cuántica", 1)


def add_section(title, paragraphs):
    doc.add_heading(title, level=2)
    for p in paragraphs:
        doc.add_paragraph(p)


def add_subsection(title, paragraphs=None, bullets=None):
    doc.add_heading(title, level=3)
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if bullets:
        for item in bullets:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(item)


def add_numbered(title, items):
    doc.add_heading(title, level=3)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for run in hdr_cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph("")


def add_code(lines):
    p = doc.add_paragraph()
    run = p.add_run(lines)
    run.font.name = "Consolas"
    run.font.size = doc.styles["Normal"].font.size
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


add_section(
    "1. Introducción y objetivos",
    [
        "En las sesiones anteriores vimos que la seguridad de RSA y ECC depende de problemas matemáticos difíciles para un ordenador clásico: factorizar números grandes y calcular logaritmos discretos. La computación cuántica es un modelo de cálculo distinto, basado en la mecánica cuántica, que resuelve estos problemas de forma radicalmente más eficiente, y por eso amenaza directamente a la criptografía actual.",
        "Al finalizar esta sesión el estudiante podrá:"
    ],
)
for item in [
    "explicar qué es un qubit y las nociones de superposición, entrelazamiento y medición;",
    "identificar las puertas cuánticas básicas y su función en un circuito;",
    "explicar con detalle matemático cómo el algoritmo de Shor rompe RSA y ECC;",
    "entender por qué la criptografía simétrica y el hash se debilitan (no se rompen) frente a Grover;",
    "situar esta amenaza en el camino hacia la criptografía post-cuántica."
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)

add_section(
    "2. Qubits y superposición",
    [
        "Un bit clásico vale 0 o 1. Un qubit se representa como |ψ⟩ = α|0⟩ + β|1⟩, con |α|² + |β|² = 1.",
        "Puede existir en una combinación de |0⟩ y |1⟩ a la vez: la superposición. Por ejemplo, la puerta Hadamard sobre |0⟩ produce H|0⟩ = (1/√2)(|0⟩ + |1⟩), de forma que P(0) = P(1) = 0.5.",
        "Con n qubits se pueden representar hasta 2ⁿ combinaciones a la vez, pero medir colapsa el resultado a uno solo: la ventaja cuántica está en diseñar la interferencia para que la respuesta correcta sea la más probable, no en \"probar todo a la vez\"."
    ],
)

add_section(
    "3. Entrelazamiento y medición",
    [
        "El entrelazamiento correlaciona dos o más qubits de forma que no pueden describirse por separado. El estado de Bell |Φ+⟩ = (1/√2)(|00⟩ + |11⟩) hace que, al medir un qubit, el resultado del otro quede determinado al instante, sin importar la distancia (aunque no permite comunicación más rápida que la luz).",
        "Medir un qubit colapsa su estado de forma irreversible, con probabilidades P(0) = |α|² y P(1) = |β|². El teorema de no clonación impide copiar exactamente un estado cuántico desconocido."
    ],
)

add_section(
    "4. Puertas cuánticas y circuitos básicos",
    [],
)
add_table(
    ["Puerta", "Efecto"],
    [
        ["X (NOT cuántico)", "Invierte |0⟩ y |1⟩"],
        ["H (Hadamard)", "Crea superposición"],
        ["Z", "Cambia la fase del estado |1⟩"],
        ["CNOT", "Invierte el 2º qubit según el 1º (crea entrelazamiento)"],
    ],
)
doc.add_paragraph(
    "Un circuito cuántico encadena puertas y termina con una medición. Aplicar H y después CNOT sobre dos qubits genera el estado de Bell descrito arriba."
)

add_section(
    "5. Computación clásica frente a computación cuántica",
    [],
)
add_table(
    ["Aspecto", "Computación clásica", "Computación cuántica"],
    [
        ["Unidad básica", "Bit (0 o 1)", "Qubit (superposición de 0 y 1)"],
        ["Determinismo", "Determinista", "Probabilístico (la medición da un resultado con cierta probabilidad)"],
        ["Copia de estados", "Trivial", "Prohibida por el teorema de no clonación"],
        ["Ventaja principal", "Universalidad y madurez tecnológica", "Muy superior en problemas concretos (factorización, búsqueda)"],
    ],
)
doc.add_paragraph(
    "No es un \"superordenador universal más rápido para todo\": su ventaja se limita a problemas con una estructura matemática concreta."
)

add_section(
    "6. De Deutsch-Jozsa a Grover",
    [],
)
for item in [
    "Deutsch-Jozsa (1992): distingue si una función f(x) es constante o equilibrada con una única evaluación, mientras que un algoritmo clásico determinista podría necesitar varias. Su relevancia es más histórica/conceptual que práctica.",
    "Grover (1996): búsqueda en una lista desordenada de N elementos con complejidad O(√N), frente a O(N) en el caso clásico (aceleración cuadrática). Aplicado a una clave de n bits, reduce el esfuerzo de 2ⁿ a 2^(n/2): por eso amenaza a AES y a las funciones hash, sin llegar a romperlos."
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)

add_section(
    "7. El algoritmo de Shor",
    [
        "El algoritmo de Shor (1994) es el más relevante para la criptografía: resuelve en tiempo polinómico la factorización de enteros y el logaritmo discreto (incluido en curvas elípticas), la base matemática de RSA, ECC y Diffie-Hellman."
    ],
)
add_subsection(
    "Idea matemática",
    paragraphs=[
        "Factorizar N = p × q se reduce a encontrar el periodo r de la función f(x) = aˣ mod N, con 1 < a < N y mcd(a, N) = 1, es decir, el menor r > 0 tal que aʳ ≡ 1 (mod N).",
        "Si r es par y a^(r/2) no es congruente con -1 módulo N, entonces:",
        "aʳ − 1 = (a^(r/2) − 1)(a^(r/2) + 1) ≡ 0 (mod N)",
        "por lo que p = mcd(a^(r/2) − 1, N) y q = mcd(a^(r/2) + 1, N) son factores no triviales de N."
    ],
)
add_numbered(
    "Pasos del algoritmo",
    [
        "Elegir a aleatorio, 1 < a < N, con mcd(a, N) = 1.",
        "(Parte cuántica) Calcular el periodo r de f(x) = aˣ mod N mediante superposición y la transformada cuántica de Fourier, en tiempo polinómico — un ordenador clásico necesitaría un esfuerzo exponencial para el mismo cálculo.",
        "Si r es impar o a^(r/2) ≡ -1 (mod N), repetir con otro a.",
        "Si no, calcular p = mcd(a^(r/2) − 1, N) y q = mcd(a^(r/2) + 1, N): son los factores de N."
    ],
)
doc.add_heading("Ejemplo numérico con N = 15, a = 7", level=3)
add_table(
    ["x", "7ˣ mod 15"],
    [
        ["1", "7"],
        ["2", "4"],
        ["3", "13"],
        ["4", "1  ← periodo r = 4 (par)"],
    ],
)
add_code(
    "a^(r/2) = 7^2 mod 15 = 4        (4 != 14 = -1 mod 15  -> cumple la condicion)\n\n"
    "p = mcd(4 - 1, 15) = mcd(3, 15) = 3\n"
    "q = mcd(4 + 1, 15) = mcd(5, 15) = 5\n\n"
    "15 = 3 x 5"
)
doc.add_paragraph(
    "Con N = 15 este cálculo también es viable a mano; la ventaja cuántica aparece cuando N tiene cientos de dígitos (como en una clave RSA real): encontrar r de forma clásica es inviable, pero el algoritmo cuántico lo hace en tiempo polinómico."
)
doc.add_paragraph(
    "Nota: hoy no existe un ordenador cuántico con suficientes qubits estables para ejecutar Shor contra claves reales, pero el riesgo (\"harvest now, decrypt later\": guardar tráfico cifrado hoy para descifrarlo en el futuro) ya obliga a planificar la migración."
)

add_section(
    "8. Impacto en la criptografía actual",
    [],
)
add_table(
    ["Algoritmo", "Base matemática", "Efecto cuántico"],
    [
        ["RSA", "Factorización de enteros", "Roto por Shor (tiempo polinómico)"],
        ["ECC / Diffie-Hellman", "Logaritmo discreto", "Roto por Shor (tiempo polinómico)"],
        ["AES", "Confusión y difusión", "Debilitado por Grover (128→64 bits efectivos); usar AES-256"],
        ["SHA-256 / SHA-3", "Funciones hash", "Debilitado por Grover; usar salidas de 384-512 bits"],
    ],
)
doc.add_paragraph(
    "En resumen: la criptografía asimétrica clásica (RSA, ECC, Diffie-Hellman) se rompe por completo; la simétrica y el hash se debilitan pero no se rompen, ajustando el tamaño de clave."
)

add_section(
    "9. Mitos y errores comunes",
    [],
)
for item in [
    "\"Prueban todas las soluciones a la vez\": la medición colapsa a un único resultado; la ventaja viene de la interferencia, no de \"probarlo todo\".",
    "\"El entrelazamiento permite comunicación instantánea\": genera correlaciones, no un canal de comunicación.",
    "\"Ya existen ordenadores cuánticos que rompen RSA\": no, aún no existen con suficientes qubits lógicos estables.",
    "\"Hace obsoleta toda la criptografía\": solo la basada en factorización/logaritmo discreto; la simétrica y el hash se ajustan aumentando el tamaño de clave."
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)

add_section(
    "10. Conclusiones y hacia la criptografía post-cuántica",
    [
        "La computación cuántica es un modelo de cálculo distinto: Shor rompe RSA/ECC/Diffie-Hellman en tiempo polinómico, mientras que Grover debilita (sin romper) AES y las funciones hash. Aunque hoy no existe un ordenador cuántico capaz de ejecutar Shor contra claves reales, el riesgo a medio/largo plazo y el \"harvest now, decrypt later\" justifican empezar ya la migración hacia la criptografía post-cuántica (PQC): algoritmos basados en retículos, códigos correctores de errores o funciones hash, actualmente en proceso de estandarización por el NIST, que se estudiarán en las próximas sesiones."
    ],
)

add_section(
    "11. Bibliografía y preguntas de reflexión",
    [],
)
add_subsection(
    "Bibliografía",
    bullets=[
        "Nielsen, M. A. y Chuang, I. L. Quantum Computation and Quantum Information. Cambridge University Press.",
        "Shor, P. W. Algorithms for quantum computation: discrete logarithms and factoring. Proceedings of the 35th Annual Symposium on Foundations of Computer Science, 1994.",
        "Grover, L. K. A fast quantum mechanical algorithm for database search. STOC, 1996.",
        "National Institute of Standards and Technology (NIST). Documentación del proceso de estandarización de criptografía post-cuántica."
    ],
)
add_numbered(
    "Preguntas de reflexión",
    [
        "¿Por qué encontrar el periodo r de aˣ mod N permite factorizar N?",
        "¿Por qué el algoritmo de Shor rompe RSA/ECC mientras que Grover solo debilita AES?",
        "¿Por qué el riesgo cuántico afecta hoy a datos cifrados aunque el ataque solo sea viable en el futuro?"
    ],
)

out_path = preferred_out_path
try:
    doc.save(out_path)
    print(f"Word generated: {out_path}")
except PermissionError:
    out_path = backup_out_path
    doc.save(out_path)
    print(f"Word generated: {out_path}")
    print("The main output file was locked by another application. A refined copy was generated instead.")
