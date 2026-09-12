from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "05_entregables" / "05_evaluacion"


def add_markdown_to_doc(doc: Document, text: str) -> None:
    lines = text.splitlines()
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            continue

        if line.startswith("**Respuesta correcta:**"):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            continue

        if line.startswith("**Retroalimentación:**"):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.italic = True
            continue

        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:].strip())
            continue

        doc.add_paragraph(line)


def generate_word_from_markdown(md_path: Path) -> Path:
    doc = Document()
    content = md_path.read_text(encoding="utf-8")
    add_markdown_to_doc(doc, content)

    out_path = md_path.with_suffix(".docx")
    doc.save(out_path)
    return out_path


def main() -> None:
    for md_file in sorted(INPUT_DIR.glob("test_*.md")):
        out_path = generate_word_from_markdown(md_file)
        print(f"Generated: {out_path.name}")


if __name__ == "__main__":
    main()
