from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT_PATH = Path(r"C:\Users\cmari\Desktop\IEP\CriptografíaPostcuantica\05_entregables\06_presentaciones\sesion_2_criptografia_asimetrica.pptx")
OUT_PATH.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

GRAY = RGBColor(242, 242, 242)
DARK = RGBColor(27, 38, 49)
BLUE = RGBColor(20, 86, 136)
BLUE_LIGHT = RGBColor(91, 149, 191)
TEXT = RGBColor(38, 50, 56)
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(17, 112, 157)
RED = RGBColor(235, 104, 92)


def add_bg_fill(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def add_header(slide, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GRAY
    bar.line.color.rgb = GRAY

    tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.20), Inches(11.5), Inches(0.45))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = DARK


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.8), Inches(1.0))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = DARK

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.0), Inches(0.5))
        p = sub_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.color.rgb = BLUE
        run.font.bold = True


def add_bullets(slide, bullets, left=0.9, top=1.7, width=11.4, height=4.8, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(size)
        run.font.color.rgb = TEXT


def add_two_columns(slide, header_title, left_items, right_items, title_left, title_right):
    add_header(slide, header_title)

    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.7), Inches(4.9))
    left_tf = left_box.text_frame
    left_tf.word_wrap = True
    p = left_tf.paragraphs[0]
    p.text = title_left
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(10)
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLUE

    for item in left_items:
        p = left_tf.add_paragraph()
        p.text = item
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(17)
        run.font.color.rgb = TEXT

    right_box = slide.shapes.add_textbox(Inches(6.9), Inches(1.5), Inches(5.7), Inches(4.9))
    right_tf = right_box.text_frame
    right_tf.word_wrap = True
    p = right_tf.paragraphs[0]
    p.text = title_right
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(10)
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLUE

    for item in right_items:
        p = right_tf.add_paragraph()
        p.text = item
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(17)
        run.font.color.rgb = TEXT


def add_table_slide(slide, header_title, subtitle, col_headers, rows, left=0.9, top=1.9, width=11.5, height=4.2):
    add_header(slide, header_title)
    add_title(slide, subtitle)

    n_rows = len(rows) + 1
    n_cols = len(col_headers)
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = graphic_frame.table

    for c, text in enumerate(col_headers):
        cell = table.cell(0, c)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name = 'Calibri'
                run.font.bold = True
                run.font.size = Pt(15)
                run.font.color.rgb = WHITE

    for r, row_values in enumerate(rows, start=1):
        for c, value in enumerate(row_values):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = GRAY if r % 2 == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(13)
                    run.font.color.rgb = TEXT


def new_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg_fill(slide)
    return slide


# ---------------------------------------------------------------------------
# Slide 1: Portada
# ---------------------------------------------------------------------------
slide = new_slide()

shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.8), Inches(1.3))
shape.fill.solid(); shape.fill.fore_color.rgb = GRAY; shape.line.color.rgb = GRAY
p = shape.text_frame.paragraphs[0]
p.text = 'Criptografía asimétrica: RSA, ECC y protocolos básicos'
p.alignment = PP_ALIGN.LEFT
run = p.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(28); run.font.bold = True; run.font.color.rgb = DARK

sub = slide.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(7.2), Inches(0.8))
sub_text = sub.text_frame.paragraphs[0]
sub_text.text = 'Sesión 2'
run = sub_text.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(20); run.font.bold = True; run.font.color.rgb = BLUE

sub2 = slide.shapes.add_textbox(Inches(0.9), Inches(3.4), Inches(9.2), Inches(1.8))
sub2_text = sub2.text_frame.paragraphs[0]
sub2_text.text = 'Claves públicas y privadas, RSA paso a paso, curvas elípticas, Diffie-Hellman y certificados digitales.'
run = sub2_text.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(18); run.font.color.rgb = TEXT

box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.6), Inches(2.2), Inches(2.5), Inches(2.5))
box.fill.solid(); box.fill.fore_color.rgb = BLUE; box.line.color.rgb = BLUE
inner = slide.shapes.add_textbox(Inches(9.9), Inches(2.75), Inches(2.0), Inches(1.2))
inner_text = inner.text_frame.paragraphs[0]
inner_text.text = '02'
inner_text.alignment = PP_ALIGN.CENTER
run = inner_text.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(28); run.font.bold = True; run.font.color.rgb = WHITE

# ---------------------------------------------------------------------------
# Slide 2: Objetivos de la sesión
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'Objetivos de la sesión')
add_title(slide, '¿Qué aprenderemos hoy?', 'De la teoría al ejemplo numérico')
add_bullets(slide, [
    'Repasar los fundamentos de la criptografía asimétrica: clave pública y clave privada.',
    'Entender cómo se genera una clave RSA paso a paso con un ejemplo numérico completo.',
    'Ver cómo se cifra, se descifra y se firma con RSA.',
    'Conocer ECC y por qué ofrece la misma seguridad con claves más pequeñas.',
    'Comprender Diffie-Hellman y su papel en el acuerdo de claves.',
    'Conectar todo con la práctica real: certificados digitales y TLS/HTTPS.',
])

# ---------------------------------------------------------------------------
# Slide 3: Conceptos clave de criptografía asimétrica
# ---------------------------------------------------------------------------
slide = new_slide()
add_two_columns(
    slide,
    'Criptografía asimétrica: conceptos clave',
    [
        'Se puede compartir libremente con cualquiera.',
        'Se usa para cifrar mensajes dirigidos a su propietario.',
        'Se usa para verificar firmas digitales.',
    ],
    [
        'Debe mantenerse siempre en secreto.',
        'Se usa para descifrar mensajes recibidos.',
        'Se usa para firmar digitalmente.',
    ],
    title_left='Clave pública',
    title_right='Clave privada',
)

t = slide.shapes.add_textbox(Inches(0.9), Inches(6.6), Inches(11.4), Inches(0.6))
p = t.text_frame.paragraphs[0]
p.text = 'Regla de oro: para cifrar, pública → privada. Para firmar, privada → pública.'
run = p.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(16); run.font.bold = True; run.font.color.rgb = ACCENT

# ---------------------------------------------------------------------------
# Slide 4: Qué es RSA
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'RSA: definición y funcionamiento')
add_title(slide, '¿Qué es RSA?', 'Rivest, Shamir y Adleman (1977)')
add_bullets(slide, [
    'Algoritmo de criptografía asimétrica usado para cifrar, firmar y autenticar.',
    'Su seguridad se basa en un hecho matemático simple: multiplicar dos primos grandes es fácil.',
    'Pero factorizar el resultado (recuperar esos dos primos) es extremadamente difícil.',
    'Esa dificultad de factorización es la base de la seguridad de RSA.',
    'Se usa junto a esquemas de relleno seguro: RSA-OAEP (cifrado) y RSA-PSS (firma).',
])

# ---------------------------------------------------------------------------
# Slide 5: Ejemplo didáctico RSA - generación de claves
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'RSA paso a paso')
add_title(slide, 'Ejemplo didáctico: generación de claves', 'Usamos números pequeños solo para entender el proceso')
add_bullets_small = add_bullets  # reutilizamos el mismo helper con tamaño reducido
add_bullets(slide, [
    '1. Elegimos dos primos: p = 5, q = 11.',
    '2. Calculamos n = p × q = 55.',
    '3. Calculamos φ(n) = (p−1)(q−1) = 4 × 10 = 40.',
    '4. Elegimos e = 3, que cumple gcd(3, 40) = 1.',
    '5. Calculamos d, el inverso modular de e: 3 × 27 = 81 ≡ 1 (mod 40) → d = 27.',
    '6. Clave pública = (e, n) = (3, 55).   Clave privada = (d, n) = (27, 55).',
], size=19)

# ---------------------------------------------------------------------------
# Slide 6: Cifrado, descifrado y firma con RSA
# ---------------------------------------------------------------------------
slide = new_slide()
add_table_slide(
    slide,
    'RSA en acción',
    'Cifrado, descifrado y firma digital',
    col_headers=['Operación', 'Clave usada', 'Fórmula', 'Propósito'],
    rows=[
        ['Cifrar', 'Pública del destinatario', 'c = m^e mod n', 'Confidencialidad'],
        ['Descifrar', 'Privada del destinatario', 'm = c^d mod n', 'Recuperar el mensaje'],
        ['Firmar', 'Privada del firmante', 's = h^d mod n', 'Autenticidad e integridad'],
        ['Verificar', 'Pública del firmante', "h' = s^e mod n", 'Comprobar la firma'],
    ],
    top=2.1,
    height=3.2,
)

t = slide.shapes.add_textbox(Inches(0.9), Inches(5.7), Inches(11.4), Inches(1.3))
tf = t.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = 'Ejemplo numérico: con m = 10, clave pública (3, 55) y privada (27, 55):'
p.runs[0].font.name = 'Calibri'; p.runs[0].font.size = Pt(16); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = TEXT
p2 = tf.add_paragraph()
p2.text = 'c = 10^3 mod 55 = 10      →      m = 10^27 mod 55 = 10 (se recupera el mensaje original)'
p2.runs[0].font.name = 'Calibri'; p2.runs[0].font.size = Pt(16); p2.runs[0].font.color.rgb = ACCENT

# ---------------------------------------------------------------------------
# Slide 7: ECC
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'ECC: criptografía de curva elíptica')
add_title(slide, '¿Qué es ECC?', 'Otra forma de hacer criptografía asimétrica')
add_bullets(slide, [
    'Se basa en el problema del logaritmo discreto sobre curvas elípticas (y² = x³ + ax + b).',
    'Conocer un punto P y Q = kP no permite recuperar fácilmente el valor k.',
    'Ofrece el mismo nivel de seguridad que RSA, pero con claves mucho más pequeñas.',
    'Ideal para móviles, tarjetas inteligentes y dispositivos IoT con recursos limitados.',
    'Es la base de ECDH (acuerdo de claves) y ECDSA/EdDSA (firma digital).',
])

# ---------------------------------------------------------------------------
# Slide 8: Comparativa RSA vs ECC
# ---------------------------------------------------------------------------
slide = new_slide()
add_table_slide(
    slide,
    'Comparativa RSA frente a ECC',
    'Misma seguridad, distinto coste',
    col_headers=['Aspecto', 'RSA', 'ECC'],
    rows=[
        ['Base matemática', 'Factorización de enteros', 'Logaritmo discreto en curvas'],
        ['Clave equivalente', '2048 bits', '~224-256 bits'],
        ['Velocidad', 'Más lenta con claves grandes', 'Más rápida y eficiente'],
        ['Consumo en IoT/móvil', 'Alto', 'Bajo'],
        ['Uso típico', 'Certificados clásicos', 'TLS moderno, IoT, móviles'],
    ],
    top=2.0,
    height=3.6,
)

# ---------------------------------------------------------------------------
# Slide 9: Diffie-Hellman
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'Diffie-Hellman: acuerdo de claves')
add_title(slide, '¿Cómo se acuerda un secreto sin enviarlo?', 'El problema que resuelve DH')
add_bullets(slide, [
    'Alice y Bob acuerdan públicamente un primo p y un generador g.',
    'Alice calcula A = g^a mod p (a es secreto); Bob calcula B = g^b mod p (b es secreto).',
    'Intercambian A y B por el canal público.',
    'Ambos llegan al mismo secreto: K = g^(ab) mod p.',
    'Limitación clave: DH no autentica por sí solo → riesgo de ataque de intermediario (MITM).',
    'Solución: combinar DH/ECDH con autenticación mediante certificados digitales.',
])

# ---------------------------------------------------------------------------
# Slide 10: Certificados, PKI y TLS
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'De la teoría a la práctica')
add_title(slide, 'Certificados digitales, PKI y TLS/HTTPS', 'Todo lo visto, funcionando junto')
add_bullets(slide, [
    'Un certificado digital vincula una clave pública con la identidad de su propietario.',
    'Una autoridad certificadora (CA) firma los certificados y sostiene la cadena de confianza.',
    'TLS combina: intercambio de claves (ECDH), autenticación (certificado) y cifrado simétrico (AES-GCM).',
    'Así se protege la mayoría de las conexiones web actuales (HTTPS).',
])

# ---------------------------------------------------------------------------
# Slide 11: Cierre
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'Cierre y puntos clave')
add_title(slide, 'Lo que debes recordar', 'Errores comunes y buenas prácticas')
add_bullets(slide, [
    'RSA: seguro gracias a la dificultad de factorizar números grandes; usar siempre con OAEP/PSS.',
    'ECC: misma seguridad que RSA con claves mucho más pequeñas; ideal para IoT y móviles.',
    'Diffie-Hellman necesita autenticación adicional para evitar ataques de intermediario.',
    'Nunca implementes tu propio esquema criptográfico: usa librerías revisadas y mantenidas.',
    'Certificados y PKI son la pieza que conecta la criptografía asimétrica con la confianza real en Internet.',
])

t = slide.shapes.add_textbox(Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.4))
p = t.text_frame.paragraphs[0]
p.text = 'Pregunta de reflexión: ¿qué pasaría si se reutilizaran los mismos primos p y q en varias claves RSA?'
p.runs[0].font.name = 'Calibri'; p.runs[0].font.size = Pt(12); p.runs[0].font.color.rgb = RGBColor(90, 90, 90)

prs.save(OUT_PATH)
print(f'PowerPoint generated: {OUT_PATH}')
