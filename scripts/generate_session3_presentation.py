from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT_PATH = Path(r"C:\Users\cmari\Desktop\IEP\CriptografíaPostcuantica\05_entregables\06_presentaciones\sesion_3_computacion_cuantica.pptx")
OUT_PATH.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

GRAY = RGBColor(242, 242, 242)
DARK = RGBColor(27, 38, 49)
BLUE = RGBColor(20, 86, 136)
BLUE_LIGHT = RGBColor(91, 149, 191)
TEXT = RGBColor(38, 50, 56)
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(17, 112, 157)
RED = RGBColor(235, 104, 92)


def add_bg_fill(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def add_header(slide, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GRAY
    bar.line.color.rgb = GRAY

    tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.20), Inches(11.5), Inches(0.45))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = DARK


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.8), Inches(1.0))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = DARK

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.0), Inches(0.5))
        p = sub_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.color.rgb = BLUE
        run.font.bold = True


def add_bullets(slide, bullets, left=0.9, top=1.7, width=11.4, height=4.8, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(size)
        run.font.color.rgb = TEXT


def add_two_columns(slide, header_title, left_items, right_items, title_left, title_right):
    add_header(slide, header_title)

    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.7), Inches(4.9))
    left_tf = left_box.text_frame
    left_tf.word_wrap = True
    p = left_tf.paragraphs[0]
    p.text = title_left
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(10)
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLUE

    for item in left_items:
        p = left_tf.add_paragraph()
        p.text = item
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(17)
        run.font.color.rgb = TEXT

    right_box = slide.shapes.add_textbox(Inches(6.9), Inches(1.5), Inches(5.7), Inches(4.9))
    right_tf = right_box.text_frame
    right_tf.word_wrap = True
    p = right_tf.paragraphs[0]
    p.text = title_right
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(10)
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLUE

    for item in right_items:
        p = right_tf.add_paragraph()
        p.text = item
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(17)
        run.font.color.rgb = TEXT


def add_table_slide(slide, header_title, subtitle, col_headers, rows, left=0.9, top=1.9, width=11.5, height=4.2):
    add_header(slide, header_title)
    add_title(slide, subtitle)

    n_rows = len(rows) + 1
    n_cols = len(col_headers)
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = graphic_frame.table

    for c, text in enumerate(col_headers):
        cell = table.cell(0, c)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name = 'Calibri'
                run.font.bold = True
                run.font.size = Pt(15)
                run.font.color.rgb = WHITE

    for r, row_values in enumerate(rows, start=1):
        for c, value in enumerate(row_values):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = GRAY if r % 2 == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(13)
                    run.font.color.rgb = TEXT


def new_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg_fill(slide)
    return slide


# ---------------------------------------------------------------------------
# Slide 1: Portada
# ---------------------------------------------------------------------------
slide = new_slide()

shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.8), Inches(1.3))
shape.fill.solid(); shape.fill.fore_color.rgb = GRAY; shape.line.color.rgb = GRAY
p = shape.text_frame.paragraphs[0]
p.text = 'Introducción a la computación cuántica'
p.alignment = PP_ALIGN.LEFT
run = p.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(28); run.font.bold = True; run.font.color.rgb = DARK

sub = slide.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(7.2), Inches(0.8))
sub_text = sub.text_frame.paragraphs[0]
sub_text.text = 'Sesión 3 — versión de 30 minutos'
run = sub_text.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(20); run.font.bold = True; run.font.color.rgb = BLUE

sub2 = slide.shapes.add_textbox(Inches(0.9), Inches(3.4), Inches(9.2), Inches(1.8))
sub2_text = sub2.text_frame.paragraphs[0]
sub2_text.text = 'Qubits, superposición, entrelazamiento y el algoritmo de Shor: por qué la computación cuántica amenaza a RSA y ECC.'
run = sub2_text.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(18); run.font.color.rgb = TEXT

box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.6), Inches(2.2), Inches(2.5), Inches(2.5))
box.fill.solid(); box.fill.fore_color.rgb = BLUE; box.line.color.rgb = BLUE
inner = slide.shapes.add_textbox(Inches(9.9), Inches(2.75), Inches(2.0), Inches(1.2))
inner_text = inner.text_frame.paragraphs[0]
inner_text.text = '03'
inner_text.alignment = PP_ALIGN.CENTER
run = inner_text.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(28); run.font.bold = True; run.font.color.rgb = WHITE

# ---------------------------------------------------------------------------
# Slide 2: Objetivos de la sesión
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'Objetivos de la sesión')
add_title(slide, '¿Qué aprenderemos hoy?', 'De los qubits al algoritmo que rompe RSA')
add_bullets(slide, [
    'Explicar qué es un qubit y las nociones de superposición, entrelazamiento y medición.',
    'Identificar las puertas cuánticas básicas y su función en un circuito.',
    'Explicar con detalle matemático cómo el algoritmo de Shor rompe RSA y ECC.',
    'Entender por qué la criptografía simétrica y el hash se debilitan (no se rompen) frente a Grover.',
    'Situar esta amenaza en el camino hacia la criptografía post-cuántica.',
])

# ---------------------------------------------------------------------------
# Slide 3: Qubits y superposición
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'Qubits y superposición')
add_title(slide, 'Un qubit puede ser 0 y 1 a la vez', '|ψ⟩ = α|0⟩ + β|1⟩,  |α|² + |β|² = 1')
add_bullets(slide, [
    'Un bit clásico solo vale 0 o 1; un qubit puede estar en superposición de |0⟩ y |1⟩.',
    'La puerta Hadamard sobre |0⟩ produce H|0⟩ = (1/√2)(|0⟩ + |1⟩), con P(0) = P(1) = 0.5.',
    'Con n qubits se representan hasta 2ⁿ combinaciones a la vez, pero medir colapsa el resultado a uno solo.',
    'La ventaja cuántica está en diseñar la interferencia para que la respuesta correcta sea la más probable, no en "probar todo a la vez".',
])

# ---------------------------------------------------------------------------
# Slide 4: Entrelazamiento y medición
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'Entrelazamiento y medición')
add_title(slide, 'Correlaciones que no existen en el mundo clásico', 'Estado de Bell: (1/√2)(|00⟩ + |11⟩)')
add_bullets(slide, [
    'Dos qubits entrelazados están correlacionados aunque estén físicamente separados.',
    'Al medir uno, el resultado del otro queda determinado al instante (sin transmitir información más rápido que la luz).',
    'Medir colapsa el estado de forma irreversible: P(0) = |α|², P(1) = |β|².',
    'El teorema de no clonación impide copiar exactamente un estado cuántico desconocido.',
])

# ---------------------------------------------------------------------------
# Slide 5: Puertas cuánticas
# ---------------------------------------------------------------------------
slide = new_slide()
add_table_slide(
    slide,
    'Puertas cuánticas básicas',
    'Los "bloques" de un circuito cuántico',
    col_headers=['Puerta', 'Efecto principal'],
    rows=[
        ['X (NOT cuántico)', 'Invierte |0⟩ y |1⟩'],
        ['H (Hadamard)', 'Crea superposición'],
        ['Z', 'Cambia la fase del estado |1⟩'],
        ['CNOT', 'Invierte el 2º qubit según el 1º; crea entrelazamiento'],
    ],
    top=2.1,
    height=3.0,
)
t = slide.shapes.add_textbox(Inches(0.9), Inches(5.4), Inches(11.4), Inches(1.2))
tf = t.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = 'Un circuito cuántico encadena puertas y termina con una medición. H + CNOT sobre dos qubits genera el estado de Bell.'
p.runs[0].font.name = 'Calibri'; p.runs[0].font.size = Pt(16); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = TEXT

# ---------------------------------------------------------------------------
# Slide 6: Computación clásica frente a cuántica
# ---------------------------------------------------------------------------
slide = new_slide()
add_table_slide(
    slide,
    'Computación clásica frente a cuántica',
    'No es "más rápido", es distinto',
    col_headers=['Aspecto', 'Clásica', 'Cuántica'],
    rows=[
        ['Unidad básica', 'Bit (0 o 1)', 'Qubit (superposición)'],
        ['Determinismo', 'Determinista', 'Probabilístico al medir'],
        ['Copia de estados', 'Trivial', 'Prohibida (no clonación)'],
        ['Ventaja principal', 'Universalidad y madurez', 'Muy superior en problemas concretos'],
    ],
    top=2.0,
    height=3.4,
)

# ---------------------------------------------------------------------------
# Slide 7: De Deutsch-Jozsa a Grover
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'De Deutsch-Jozsa a Grover')
add_title(slide, 'Los primeros algoritmos con ventaja cuántica', 'Antes de llegar a Shor')
add_bullets(slide, [
    'Deutsch-Jozsa (1992): distingue si f(x) es constante o equilibrada con una única evaluación (clásicamente puede necesitar varias). Relevancia más histórica que práctica.',
    'Grover (1996): búsqueda en una lista desordenada de N elementos en O(√N), frente a O(N) clásico (aceleración cuadrática).',
    'Aplicado a una clave de n bits, Grover reduce el esfuerzo de 2ⁿ a 2^(n/2): amenaza a AES y al hash, sin llegar a romperlos.',
])

# ---------------------------------------------------------------------------
# Slide 8: El algoritmo de Shor — idea matemática
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'El algoritmo de Shor: la idea matemática')
add_title(slide, 'Factorizar N = p × q, resuelto en tiempo polinómico', 'La base de RSA, ECC y Diffie-Hellman se rompe')
add_bullets(slide, [
    'Factorizar N se reduce a encontrar el periodo r de f(x) = aˣ mod N, con 1 < a < N y mcd(a, N) = 1: el menor r > 0 tal que aʳ ≡ 1 (mod N).',
    'Si r es par y a^(r/2) no es congruente con -1 mod N: aʳ − 1 = (a^(r/2) − 1)(a^(r/2) + 1) ≡ 0 (mod N).',
    'Entonces p = mcd(a^(r/2) − 1, N) y q = mcd(a^(r/2) + 1, N) son factores no triviales de N.',
    '(Parte cuántica) Encontrar r mediante superposición y la transformada cuántica de Fourier es polinómico; clásicamente sería exponencial.',
], size=19)

# ---------------------------------------------------------------------------
# Slide 9: El algoritmo de Shor — ejemplo numérico
# ---------------------------------------------------------------------------
slide = new_slide()
add_table_slide(
    slide,
    'El algoritmo de Shor: ejemplo numérico',
    'N = 15, a = 7 → factores 3 y 5',
    col_headers=['x', '7ˣ mod 15'],
    rows=[
        ['1', '7'],
        ['2', '4'],
        ['3', '13'],
        ['4', '1   ← periodo r = 4 (par)'],
    ],
    top=2.0,
    height=2.6,
    width=6.0,
)
t = slide.shapes.add_textbox(Inches(7.3), Inches(2.0), Inches(5.2), Inches(3.2))
tf = t.text_frame
tf.word_wrap = True
lines = [
    ('a^(r/2) = 7² mod 15 = 4', True),
    ('(4 ≠ 14 ≡ −1 mod 15 → cumple la condición)', False),
    ('', False),
    ('p = mcd(4−1, 15) = mcd(3, 15) = 3', True),
    ('q = mcd(4+1, 15) = mcd(5, 15) = 5', True),
    ('', False),
    ('15 = 3 × 5', True),
]
for idx, (text, bold) in enumerate(lines):
    para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    para.text = text
    para.space_after = Pt(8)
    if text:
        run = para.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.bold = bold
        run.font.color.rgb = ACCENT if bold else TEXT

t2 = slide.shapes.add_textbox(Inches(0.9), Inches(5.1), Inches(11.4), Inches(1.6))
tf2 = t2.text_frame
tf2.word_wrap = True
p = tf2.paragraphs[0]
p.text = 'Con N = 15 el cálculo es viable a mano; con un N de cientos de dígitos (clave RSA real) encontrar r es inviable de forma clásica, pero el algoritmo cuántico lo hace en tiempo polinómico.'
p.runs[0].font.name = 'Calibri'; p.runs[0].font.size = Pt(15); p.runs[0].font.color.rgb = TEXT
p2 = tf2.add_paragraph()
p2.text = 'Hoy no existe un ordenador cuántico con suficientes qubits estables para ejecutar Shor contra claves reales, pero el riesgo ("harvest now, decrypt later") ya obliga a planificar la migración.'
p2.space_before = Pt(6)
p2.runs[0].font.name = 'Calibri'; p2.runs[0].font.size = Pt(15); p2.runs[0].font.bold = True; p2.runs[0].font.color.rgb = RED

# ---------------------------------------------------------------------------
# Slide 10: Impacto en la criptografía actual
# ---------------------------------------------------------------------------
slide = new_slide()
add_table_slide(
    slide,
    'Impacto en la criptografía actual',
    'Qué se rompe y qué solo se debilita',
    col_headers=['Algoritmo', 'Base matemática', 'Efecto cuántico'],
    rows=[
        ['RSA', 'Factorización de enteros', 'Roto por Shor (tiempo polinómico)'],
        ['ECC / Diffie-Hellman', 'Logaritmo discreto', 'Roto por Shor (tiempo polinómico)'],
        ['AES', 'Confusión y difusión', 'Debilitado por Grover (128→64 bits); usar AES-256'],
        ['SHA-256 / SHA-3', 'Funciones hash', 'Debilitado por Grover; usar salidas de 384-512 bits'],
    ],
    top=2.0,
    height=3.4,
)

# ---------------------------------------------------------------------------
# Slide 11: Mitos y errores comunes
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'Mitos y errores comunes')
add_title(slide, 'Lo que NO es la computación cuántica', 'Separando ciencia de exageración')
add_bullets(slide, [
    'No "prueba todas las soluciones a la vez": medir colapsa el resultado a uno solo.',
    'El entrelazamiento no permite comunicación instantánea ni más rápida que la luz.',
    'Hoy no existe ningún ordenador cuántico capaz de romper RSA o ECC reales.',
    'No hace obsoleta toda la criptografía: solo la basada en factorización y logaritmo discreto.',
])

# ---------------------------------------------------------------------------
# Slide 12: Cierre
# ---------------------------------------------------------------------------
slide = new_slide()
add_header(slide, 'Cierre y puntos clave')
add_title(slide, 'Hacia la criptografía post-cuántica', 'Lo que debes recordar de esta sesión')
add_bullets(slide, [
    'Shor rompe RSA, ECC y Diffie-Hellman en tiempo polinómico; Grover debilita (sin romper) AES y el hash.',
    'Hoy no existe un ordenador cuántico capaz de ejecutar Shor contra claves reales, pero el riesgo ya es real por el "harvest now, decrypt later".',
    'La criptografía post-cuántica (retículos, códigos, hash-based, entre otros) es la respuesta a esta amenaza.',
    'Las próximas sesiones profundizan en estos algoritmos post-cuánticos y su estandarización por el NIST.',
])

t = slide.shapes.add_textbox(Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.4))
p = t.text_frame.paragraphs[0]
p.text = 'Pregunta de reflexión: ¿por qué encontrar el periodo r de aˣ mod N permite factorizar N?'
p.runs[0].font.name = 'Calibri'; p.runs[0].font.size = Pt(12); p.runs[0].font.color.rgb = RGBColor(90, 90, 90)

prs.save(OUT_PATH)
print(f'PowerPoint generated: {OUT_PATH}')
