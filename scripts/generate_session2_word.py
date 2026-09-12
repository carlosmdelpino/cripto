from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

preferred_out_path = Path("c:/Users/cmari/Desktop/IEP/CriptografíaPostcuantica/05_entregables/02_word/sesion_2_criptografia_asimetrica.docx")
backup_out_path = preferred_out_path.with_name("sesion_2_criptografia_asimetrica_v2.docx")
preferred_out_path.parent.mkdir(parents=True, exist_ok=True)

doc = Document()

doc.add_heading("CRIPTOGRAFÍA ASIMÉTRICA: RSA, ECC Y PROTOCOLOS BÁSICOS", 0)
doc.add_paragraph("TEMA 2")
doc.add_heading("Criptografía asimétrica", 1)


def add_section(title, paragraphs):
    doc.add_heading(title, level=2)
    for p in paragraphs:
        doc.add_paragraph(p)


def add_subsection(title, paragraphs=None, bullets=None):
    doc.add_heading(title, level=3)
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if bullets:
        for item in bullets:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(item)


def add_numbered(title, items):
    doc.add_heading(title, level=3)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for run in hdr_cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph("")


def add_code(lines):
    p = doc.add_paragraph()
    run = p.add_run(lines)
    run.font.name = "Consolas"
    run.font.size = doc.styles["Normal"].font.size
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


add_section(
    "1. Introducción y repaso de la sesión anterior",
    [
        "En la sesión 1 vimos los fundamentos de la criptografía: el modelo de amenaza (Alice, Bob y Eve), los objetivos de seguridad (confidencialidad, integridad, autenticación y no repudio) y las primitivas criptográficas básicas (cifrado, hash, MAC, firma digital e intercambio de claves).",
        "También distinguimos entre criptografía simétrica, donde emisor y receptor comparten la misma clave, y criptografía asimétrica, donde cada parte dispone de un par de claves: una pública y una privada.",
        "En esta sesión profundizamos en la criptografía asimétrica, centrándonos en dos algoritmos fundamentales: RSA y ECC (Elliptic Curve Cryptography). Estos dos sistemas son la base de la mayoría de los protocolos de seguridad que usamos a diario, como HTTPS, la firma de documentos electrónicos o la autenticación de servidores."
    ],
)

add_section(
    "2. Objetivos de la sesión",
    [
        "Al finalizar esta sesión, el estudiante será capaz de:",
        "Explicar qué es la criptografía asimétrica y en qué se diferencia de la simétrica.",
        "Describir el funcionamiento de RSA, incluyendo la generación de claves, el cifrado y la firma digital.",
        "Realizar un ejemplo numérico completo de RSA con números pequeños.",
        "Explicar qué es ECC y por qué ofrece ventajas frente a RSA.",
        "Comparar RSA y ECC en términos de seguridad, tamaño de clave y rendimiento.",
        "Entender el protocolo Diffie-Hellman y su papel en el acuerdo de claves.",
        "Relacionar estos conceptos con protocolos reales como TLS/HTTPS y los certificados digitales."
    ],
)

add_section(
    "3. Criptografía asimétrica: conceptos clave",
    [
        "La criptografía asimétrica utiliza un par de claves: una clave pública, que puede compartirse libremente con cualquiera, y una clave privada, que debe mantenerse en secreto y solo la conoce su propietario.",
        "Esto resuelve uno de los grandes problemas de la criptografía simétrica: la necesidad de distribuir una clave secreta compartida antes de poder comunicarse de forma segura.",
        "La lógica de uso depende de la finalidad: para cifrado, la clave pública se usa para cifrar y la clave privada se usa para descifrar. Para firma digital, la clave privada se usa para firmar y la clave pública se usa para verificar.",
        "Estas dos reglas son la base de todo lo que veremos en RSA y ECC durante esta sesión."
    ],
)

add_section(
    "4. RSA: definición y funcionamiento",
    [
        "RSA (Rivest, Shamir y Adleman, 1977) es el algoritmo de criptografía asimétrica más conocido y uno de los primeros en implementarse de forma práctica. Se puede usar para cifrar datos, intercambiar secretos de forma segura, firmar digitalmente documentos y mensajes, y autenticar el origen de la información.",
        "Su seguridad se basa en una idea matemática muy simple de enunciar, pero muy difícil de romper en la práctica: es fácil multiplicar dos números primos grandes, pero es muy difícil factorizar el resultado si esos primos son suficientemente grandes.",
        "Es decir, la seguridad de RSA descansa en la dificultad computacional de la factorización de números enteros grandes."
    ],
)

add_subsection(
    "4.1 Las dos claves de RSA",
    bullets=[
        "Clave pública: se puede compartir con cualquiera.",
        "Clave privada: debe mantenerse en secreto.",
        "Para cifrado: pública → cifrar, privada → descifrar.",
        "Para firma digital: privada → firmar, pública → verificar."
    ],
)

add_section(
    "5. Fundamentos matemáticos de RSA",
    [
        "RSA se apoya en conceptos de teoría de números que ya introdujimos en la sesión 1: aritmética modular, números primos y el inverso modular."
    ],
)

add_subsection(
    "5.1 Números primos y producto",
    paragraphs=[
        "Se eligen dos números primos grandes p y q, y se calcula n = p × q. El valor n formará parte tanto de la clave pública como de la clave privada."
    ],
)

add_subsection(
    "5.2 Función φ de Euler",
    paragraphs=[
        "A continuación se calcula la función indicatriz de Euler de n: φ(n) = (p − 1)(q − 1). Este valor cuenta cuántos números menores que n son primos relativos con n, y es la pieza clave para construir el exponente privado."
    ],
)

add_subsection(
    "5.3 Elección del exponente público",
    paragraphs=[
        "Se elige un número e tal que 1 < e < φ(n) y gcd(e, φ(n)) = 1. Es decir, e debe ser primo relativo con φ(n). En la práctica se suele usar un valor fijo y estandarizado, como e = 65537."
    ],
)

add_subsection(
    "5.4 Cálculo del inverso modular",
    paragraphs=[
        "Finalmente, se calcula d, el inverso modular de e módulo φ(n): e × d ≡ 1 (mod φ(n)). El valor d es el exponente privado y solo lo conoce el propietario de la clave."
    ],
)

add_subsection(
    "5.5 Composición de las claves",
    bullets=[
        "Clave pública: (e, n)",
        "Clave privada: (d, n)"
    ],
)

add_section(
    "6. Generación de claves RSA: ejemplo didáctico",
    [
        "Vamos a construir un ejemplo completo con números pequeños para entender el proceso paso a paso. Importante: estos valores son solo para fines educativos; en un sistema real, p y q tienen cientos de dígitos."
    ],
)

add_numbered(
    "Pasos del ejemplo",
    [
        "Elegir dos números primos: p = 5, q = 11.",
        "Calcular n = p × q = 5 × 11 = 55.",
        "Calcular φ(n) = (p − 1)(q − 1) = 4 × 10 = 40.",
        "Elegir e = 3, que cumple gcd(3, 40) = 1.",
        "Calcular d, el inverso modular de e módulo 40: 3 × 27 = 81 = 2 × 40 + 1 ≡ 1 (mod 40). Por tanto, d = 27.",
        "Resultado final: clave pública (e, n) = (3, 55); clave privada (d, n) = (27, 55)."
    ],
)

add_section(
    "7. Cifrado y descifrado con RSA",
    [
        "Con las claves generadas en el apartado anterior, veamos cómo se cifra y descifra un mensaje."
    ],
)

add_subsection(
    "7.1 Fórmulas generales",
    paragraphs=[
        "Para cifrar un mensaje m (un número entero menor que n): c = m^e mod n.",
        "Para descifrar el texto cifrado c y recuperar el mensaje original: m = c^d mod n."
    ],
)

add_subsection(
    "7.2 Ejemplo numérico completo",
    paragraphs=[
        "Supongamos que Alice quiere enviar a Bob el mensaje m = 10, usando la clave pública de Bob (e, n) = (3, 55).",
        "Cifrado: c = 10^3 mod 55 = 1000 mod 55 = 10.",
        "Descifrado (Bob usa su clave privada (d, n) = (27, 55)): m = 10^27 mod 55 = 10.",
        "Bob recupera correctamente el mensaje original m = 10.",
        "Nota didáctica: en este ejemplo el resultado del cifrado coincide numéricamente con el mensaje original debido al tamaño pequeño de los números elegidos. Con primos grandes esta coincidencia no ocurre. Lo importante es entender el procedimiento: elevar al exponente correspondiente y tomar el resto módulo n.",
        "En sistemas reales, RSA no se usa nunca en crudo como en este ejemplo: se combina con esquemas de relleno seguro como RSA-OAEP para cifrado y RSA-PSS para firmas, que evitan ataques conocidos contra la versión matemática pura del algoritmo."
    ],
)

add_section(
    "8. Firma digital con RSA",
    [
        "RSA también permite firmar digitalmente un mensaje, invirtiendo el papel de las claves.",
        "1. El firmante calcula un resumen (hash) del mensaje, h = H(m).",
        "2. El firmante cifra ese resumen con su clave privada: s = h^d mod n.",
        "3. Cualquiera puede verificar la firma usando la clave pública del firmante: h' = s^e mod n.",
        "4. Si h' = h, la firma es válida: el mensaje procede del propietario de la clave privada y no ha sido modificado.",
        "En la práctica se usa el esquema RSA-PSS, que añade aleatoriedad y protecciones adicionales frente a ataques matemáticos sobre el esquema básico."
    ],
)

add_table(
    headers=["Operación", "Clave usada", "Propósito"],
    rows=[
        ["Cifrar", "Pública del destinatario", "Confidencialidad"],
        ["Descifrar", "Privada del destinatario", "Recuperar el mensaje"],
        ["Firmar", "Privada del firmante", "Autenticidad e integridad"],
        ["Verificar", "Pública del firmante", "Comprobar la firma"],
    ],
)

add_section(
    "9. ECC: criptografía de curva elíptica",
    [
        "ECC (Elliptic Curve Cryptography) es otra familia de algoritmos de criptografía asimétrica. En lugar de basar su seguridad en la dificultad de factorizar números grandes, ECC se apoya en el problema del logaritmo discreto sobre curvas elípticas, un problema matemático distinto y también muy difícil de resolver con los recursos computacionales actuales.",
        "Una curva elíptica se describe habitualmente con una ecuación del tipo y^2 = x^3 + a·x + b.",
        "Sobre esta curva se define una operación de 'suma de puntos' con propiedades matemáticas especiales. La seguridad de ECC se basa en que, conocidos un punto P y el resultado de sumarlo consigo mismo k veces (Q = kP), es computacionalmente muy difícil recuperar el valor k a partir de P y Q."
    ],
)

add_subsection(
    "Ventajas principales de ECC",
    bullets=[
        "Ofrece un nivel de seguridad equivalente al de RSA, pero con claves mucho más pequeñas.",
        "Es más eficiente en dispositivos con recursos limitados: móviles, tarjetas inteligentes, sensores IoT.",
        "Reduce el tamaño de los certificados y el tráfico en protocolos como TLS.",
        "Es la base de esquemas modernos como ECDSA (firma digital) y ECDH (acuerdo de claves)."
    ],
)

add_section(
    "10. Comparativa RSA frente a ECC",
    [
        "La siguiente tabla resume las diferencias principales entre ambos sistemas:"
    ],
)

add_table(
    headers=["Aspecto", "RSA", "ECC"],
    rows=[
        ["Base matemática", "Factorización de enteros", "Logaritmo discreto en curvas elípticas"],
        ["Tamaño de clave equivalente", "2048 bits", "~224-256 bits"],
        ["Velocidad de operación", "Más lenta con claves grandes", "Más rápida y eficiente"],
        ["Consumo en móviles/IoT", "Alto", "Bajo"],
        ["Uso típico", "Sistemas heredados, certificados clásicos", "TLS moderno, móviles, IoT, criptomonedas"],
        ["Esquemas asociados", "RSA-OAEP, RSA-PSS", "ECDH, ECDSA, EdDSA"],
    ],
)

add_section(
    "Ejemplo de equivalencia de seguridad",
    [
        "Una clave RSA de 2048 bits ofrece un nivel de seguridad comparable al de una clave ECC de aproximadamente 224-256 bits. Esto significa que ECC puede alcanzar el mismo nivel de protección con claves mucho más pequeñas, lo que se traduce en menor consumo de CPU, memoria y ancho de banda."
    ],
)

add_section(
    "11. Diffie-Hellman: acuerdo de claves",
    [
        "Diffie-Hellman (DH) es un protocolo que permite a dos partes acordar un secreto compartido sin enviarlo directamente por el canal de comunicación."
    ],
)

add_numbered(
    "Pasos del protocolo",
    [
        "Alice y Bob acuerdan públicamente un número primo p y un generador g.",
        "Alice elige un valor secreto a y calcula A = g^a mod p.",
        "Bob elige un valor secreto b y calcula B = g^b mod p.",
        "Alice y Bob se intercambian A y B por el canal (público).",
        "Alice calcula K = B^a mod p y Bob calcula K = A^b mod p.",
        "Ambos obtienen el mismo secreto compartido: K = g^(ab) mod p."
    ],
)

add_section(
    "Limitación importante",
    [
        "Diffie-Hellman por sí solo no autentica a las partes. Sin autenticación adicional, un atacante puede realizar un ataque de intermediario (MITM), negociando un secreto distinto con cada parte y haciéndose pasar por la otra. Por ello, en la práctica, DH siempre se combina con autenticación mediante certificados digitales o firmas.",
        "La variante basada en curvas elípticas, ECDH, sigue la misma lógica pero opera sobre puntos de una curva elíptica en lugar de sobre potencias modulares."
    ],
)

add_section(
    "12. Certificados digitales y PKI",
    [
        "Un certificado digital es un documento electrónico que vincula una clave pública con la identidad de su propietario (una persona, una organización o un servidor). Está firmado digitalmente por una autoridad certificadora (CA) en la que confían las partes.",
        "La infraestructura de clave pública (PKI) es el conjunto de roles, políticas y procedimientos necesarios para crear, distribuir, gestionar y revocar certificados digitales."
    ],
)

add_subsection(
    "Elementos principales de una PKI",
    bullets=[
        "Autoridad certificadora (CA): emite y firma certificados.",
        "Certificado digital: contiene la clave pública, la identidad del titular y la firma de la CA.",
        "Cadena de confianza: permite verificar un certificado remontando hasta una CA raíz de confianza.",
        "Revocación: mecanismos como CRL u OCSP para invalidar certificados comprometidos antes de su expiración."
    ],
)

add_section(
    "13. TLS/HTTPS: aplicación práctica",
    [
        "TLS (Transport Layer Security) es el protocolo que protege la mayoría de las comunicaciones web modernas (HTTPS). Combina de forma práctica todo lo visto en esta sesión:",
        "1. El cliente y el servidor negocian los algoritmos que van a usar.",
        "2. Se realiza un intercambio de claves (habitualmente basado en ECDH) para acordar un secreto compartido.",
        "3. El servidor se autentica mediante su certificado digital, firmado por una CA de confianza.",
        "4. A partir del secreto compartido se derivan claves simétricas para cifrar el resto de la comunicación con un esquema AEAD como AES-GCM o ChaCha20-Poly1305.",
        "De esta forma, TLS combina criptografía asimétrica (para autenticación e intercambio de claves), criptografía simétrica (para cifrar el tráfico de forma eficiente) y funciones hash (para verificar la integridad)."
    ],
)

add_section(
    "14. Ejemplos prácticos",
    [
        "Ejemplo 1: Alice cifra un mensaje con la clave pública de Bob. Solo Bob, con su clave privada, puede descifrarlo. Esto protege la confidencialidad.",
        "Ejemplo 2: Alice firma un documento con su clave privada. Cualquiera puede verificar la firma con la clave pública de Alice, comprobando autenticidad e integridad.",
        "Ejemplo 3: Un navegador se conecta a un servidor mediante HTTPS. El servidor presenta su certificado digital, el navegador lo valida contra una CA de confianza y ambos negocian claves de sesión mediante ECDH.",
        "Ejemplo 4: Un servicio IoT con recursos limitados utiliza ECC en lugar de RSA para reducir el consumo de batería y memoria manteniendo el mismo nivel de seguridad."
    ],
)

add_section(
    "15. Errores comunes y buenas prácticas",
    [
        "Errores comunes:",
        "Usar RSA sin relleno seguro (sin OAEP/PSS), vulnerable a ataques matemáticos conocidos.",
        "Reutilizar los mismos números primos o parámetros entre distintas claves.",
        "Usar tamaños de clave demasiado pequeños (por ejemplo, RSA de 512 o 1024 bits).",
        "No autenticar un intercambio Diffie-Hellman, exponiéndose a ataques MITM.",
        "Confiar en certificados autofirmados sin validación adecuada en entornos de producción.",
        "Buenas prácticas:",
        "Usar bibliotecas criptográficas revisadas y mantenidas, nunca implementaciones propias.",
        "Preferir ECC frente a RSA en nuevos sistemas cuando sea posible, por su eficiencia.",
        "Mantener actualizado el tamaño mínimo de clave recomendado por organismos como NIST.",
        "Validar siempre la cadena de certificados y comprobar mecanismos de revocación.",
        "Combinar intercambio de claves con autenticación (certificados o firmas)."
    ],
)

add_section(
    "16. Ejercicios propuestos con soluciones",
    [
        "Ejercicio 1: Con p = 3 y q = 11, calcula n y φ(n).",
        "Solución orientativa: n = 33, φ(n) = (3-1)(11-1) = 2 × 10 = 20.",
        "Ejercicio 2: ¿Por qué RSA no se usa nunca sin un esquema de relleno como OAEP o PSS?",
        "Solución orientativa: porque la versión pura del algoritmo es determinista y vulnerable a varios ataques matemáticos conocidos; el relleno añade aleatoriedad y estructura que los evita.",
        "Ejercicio 3: ¿Qué ventaja principal ofrece ECC frente a RSA?",
        "Solución orientativa: el mismo nivel de seguridad con claves mucho más pequeñas, lo que mejora el rendimiento y reduce el consumo de recursos.",
        "Ejercicio 4: ¿Por qué Diffie-Hellman necesita combinarse con autenticación?",
        "Solución orientativa: porque por sí solo no verifica la identidad de las partes, lo que permite un ataque de intermediario (MITM).",
        "Ejercicio 5: En una conexión HTTPS, ¿qué papel juega el certificado digital del servidor?",
        "Solución orientativa: vincula la clave pública del servidor con su identidad y permite al cliente autenticar al servidor antes de establecer el canal cifrado."
    ],
)

add_section(
    "17. Python como apoyo didáctico",
    [
        "Python permite comprobar de forma directa los cálculos de RSA vistos en esta sesión:"
    ],
)

add_code(
    "p, q = 5, 11\n"
    "n = p * q\n"
    "phi = (p - 1) * (q - 1)\n"
    "e = 3\n"
    "d = pow(e, -1, phi)  # inverso modular de e mod phi\n\n"
    "print(\"n =\", n)\n"
    "print(\"phi(n) =\", phi)\n"
    "print(\"clave publica  =\", (e, n))\n"
    "print(\"clave privada  =\", (d, n))\n\n"
    "m = 10\n"
    "c = pow(m, e, n)   # cifrado: c = m^e mod n\n"
    "m2 = pow(c, d, n)  # descifrado: m = c^d mod n\n\n"
    "print(\"mensaje original:\", m)\n"
    "print(\"mensaje cifrado :\", c)\n"
    "print(\"mensaje recuperado:\", m2)"
)

add_section(
    "Conclusión del ejercicio",
    [
        "Este ejercicio permite comprobar en segundos que el mensaje recuperado coincide con el original, reforzando de forma práctica los conceptos matemáticos vistos en la sesión."
    ],
)

add_section(
    "18. Conclusiones",
    [
        "La criptografía asimétrica resuelve un problema fundamental de la criptografía simétrica: la distribución segura de claves. RSA fue el primer gran sistema asimétrico ampliamente adoptado y sigue siendo relevante hoy, mientras que ECC ha ganado terreno gracias a su eficiencia y a que ofrece el mismo nivel de seguridad con claves mucho más pequeñas.",
        "En los sistemas reales, RSA y ECC rara vez se usan solos: se combinan con funciones hash, esquemas de relleno seguro, protocolos de acuerdo de claves como Diffie-Hellman/ECDH y certificados digitales dentro de una infraestructura PKI, todo ello orquestado por protocolos como TLS para proteger las comunicaciones en Internet."
    ],
)

add_section(
    "19. Bibliografía esencial",
    [
        "Rivest, R., Shamir, A. y Adleman, L. A Method for Obtaining Digital Signatures and Public-Key Cryptosystems. Communications of the ACM, 1978.",
        "Stallings, W. Cryptography and Network Security: Principles and Practice. Pearson.",
        "Katz, J. y Lindell, Y. Introduction to Modern Cryptography. Chapman & Hall/CRC.",
        "Menezes, A., van Oorschot, P. y Vanstone, S. Handbook of Applied Cryptography. CRC Press."
    ],
)

add_section(
    "20. Bibliografía recomendada",
    [
        "Hankerson, D., Menezes, A. y Vanstone, S. Guide to Elliptic Curve Cryptography. Springer.",
        "Ferguson, N., Schneier, B. y Kohno, T. Cryptography Engineering. Wiley.",
        "National Institute of Standards and Technology (NIST). Recommendation for Key Management (SP 800-57) y documentación sobre curvas elípticas aprobadas.",
        "Documentación oficial de TLS 1.3, X.509 y PKIX."
    ],
)

add_section(
    "21. Preguntas de reflexión",
    [
        "¿Por qué es fácil multiplicar dos primos grandes pero muy difícil factorizar su producto?",
        "¿Qué ocurriría si se reutilizaran los mismos primos p y q en varias claves RSA?",
        "¿Por qué ECC resulta más adecuado que RSA para dispositivos IoT con batería limitada?",
        "¿Qué diferencia hay entre cifrar con RSA y firmar con RSA en cuanto al uso de las claves?",
        "¿Qué papel juega la autoridad certificadora en la confianza de un certificado digital?",
        "¿Qué pasaría en una conexión HTTPS si el certificado del servidor no fuera válido y el navegador lo ignorara?"
    ],
)

out_path = preferred_out_path
try:
    doc.save(out_path)
    print(f"Word generated: {out_path}")
except PermissionError:
    out_path = backup_out_path
    doc.save(out_path)
    print(f"Word generated: {out_path}")
    print("The main output file was locked by another application. A refined copy was generated instead.")
