from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT_PATH = Path(r"C:\Users\cmari\Desktop\IEP\CriptografíaPostcuantica\05_entregables\06_presentaciones\sesion_1_fundamentos_criptografia.pptx")
OUT_PATH.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

GRAY = RGBColor(242, 242, 242)
DARK = RGBColor(27, 38, 49)
BLUE = RGBColor(20, 86, 136)
BLUE_LIGHT = RGBColor(91, 149, 191)
TEXT = RGBColor(38, 50, 56)
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(17, 112, 157)


def add_bg_fill(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def add_header(slide, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GRAY
    bar.line.color.rgb = GRAY

    tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.20), Inches(11.5), Inches(0.45))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = DARK


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.8), Inches(1.0))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = DARK

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.0), Inches(0.5))
        p = sub_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.color.rgb = BLUE
        run.font.bold = True


def add_bullets(slide, bullets, left=0.9, top=1.7, width=11.4, height=4.8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(20)
        run.font.color.rgb = TEXT


def add_bullets_small(slide, bullets, left=0.9, top=1.6, width=11.6, height=4.8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.color.rgb = TEXT


def add_two_columns(slide, left_items, right_items, title_left=' ', title_right=' '):
    add_header(slide, 'Fundamentos de la criptografía')

    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.5), Inches(4.8))
    left_tf = left_box.text_frame
    left_tf.word_wrap = True
    p = left_tf.paragraphs[0]
    p.text = title_left
    p.level = 0
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(10)
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLUE

    for item in left_items:
        p = left_tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.color.rgb = TEXT

    right_box = slide.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.5), Inches(4.8))
    right_tf = right_box.text_frame
    right_tf.word_wrap = True
    p = right_tf.paragraphs[0]
    p.text = title_right
    p.level = 0
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(10)
    run = p.runs[0]
    run.font.name = 'Calibri'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLUE

    for item in right_items:
        p = right_tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.color.rgb = TEXT


# Slide 1: portada
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg_fill(slide)

# big title
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.8), Inches(1.3))
shape.fill.solid(); shape.fill.fore_color.rgb = GRAY; shape.line.color.rgb = GRAY
text = shape.text_frame
p = text.paragraphs[0]
p.text = 'Fundamentos de la criptografía'
p.alignment = PP_ALIGN.LEFT
run = p.runs[0]
run.font.name = 'Calibri'
run.font.size = Pt(30)
run.font.bold = True
run.font.color.rgb = DARK

sub = slide.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(7.2), Inches(0.8))
sub_text = sub.text_frame.paragraphs[0]
sub_text.text = 'Sesión 1'
run = sub_text.runs[0]
run.font.name = 'Calibri'
run.font.size = Pt(20)
run.font.bold = True
run.font.color.rgb = BLUE

sub2 = slide.shapes.add_textbox(Inches(0.9), Inches(3.4), Inches(9.0), Inches(1.8))
sub2_text = sub2.text_frame.paragraphs[0]
sub2_text.text = 'Modelos de amenaza, objetivos de seguridad, criptografía simétrica y asimétrica, RSA, ECC y ejemplos prácticos.'
run = sub2_text.runs[0]
run.font.name = 'Calibri'
run.font.size = Pt(18)
run.font.color.rgb = TEXT

# accent box
box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.6), Inches(2.2), Inches(2.5), Inches(2.5))
box.fill.solid(); box.fill.fore_color.rgb = BLUE
box.line.color.rgb = BLUE

inner = slide.shapes.add_textbox(Inches(9.9), Inches(2.75), Inches(2.0), Inches(1.2))
inner_text = inner.text_frame.paragraphs[0]
inner_text.text = '01'
inner_text.alignment = PP_ALIGN.CENTER
run = inner_text.runs[0]
run.font.name = 'Calibri'; run.font.size = Pt(28); run.font.bold = True; run.font.color.rgb = WHITE

# Slide 2: objetivos
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg_fill(slide)
add_header(slide, 'Objetivos de la sesión')
add_title(slide, '¿Qué aprenderemos?', 'Al finalizar esta sesión...')
add_bullets(slide, [
    'Comprender qué es la criptografía y por qué es esencial en la seguridad digital.',
    'Identificar los objetivos básicos de seguridad: confidencialidad, integridad y autenticación.',
    'Distinguir entre criptografía simétrica, asimétrica e híbrida.',
    'Reconocer conceptos clave como clave, hash, cifrado, firma digital y MAC.',
    'Conectar la teoría con ejemplos prácticos y fundamentos matemáticos simples.',
])

# Slide 3: modelo de amenaza
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg_fill(slide)
add_header(slide, 'Modelo de amenaza')
add_title(slide, 'Alice, Bob y Eve', 'El canal no es seguro')

# diagram box
shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), Inches(2.0), Inches(1.6), Inches(1.2))
shape.fill.solid(); shape.fill.fore_color.rgb = BLUE_LIGHT; shape.line.color.rgb = BLUE_LIGHT
p = shape.text_frame.paragraphs[0]; p.text = 'Alice'; p.alignment = PP_ALIGN.CENTER; p.runs[0].font.name='Calibri'; p.runs[0].font.size=Pt(18); p.runs[0].font.bold=True; p.runs[0].font.color.rgb = WHITE

shape2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.6), Inches(2.0), Inches(1.6), Inches(1.2))
shape2.fill.solid(); shape2.fill.fore_color.rgb = BLUE; shape2.line.color.rgb = BLUE
p2 = shape2.text_frame.paragraphs[0]; p2.text='Bob'; p2.alignment=PP_ALIGN.CENTER; p2.runs[0].font.name='Calibri'; p2.runs[0].font.size=Pt(18); p2.runs[0].font.bold=True; p2.runs[0].font.color.rgb=WHITE

shape3 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.3), Inches(3.7), Inches(1.7), Inches(1.2))
shape3.fill.solid(); shape3.fill.fore_color.rgb = RGBColor(235, 104, 92); shape3.line.color.rgb = RGBColor(235,104,92)
p3 = shape3.text_frame.paragraphs[0]; p3.text='Eve'; p3.alignment=PP_ALIGN.CENTER; p3.runs[0].font.name='Calibri'; p3.runs[0].font.size=Pt(18); p3.runs[0].font.bold=True; p3.runs[0].font.color.rgb=WHITE

line1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(2.3), Inches(3.2), Inches(0.08))
line1.fill.solid(); line1.fill.fore_color.rgb = DARK; line1.line.color.rgb = DARK
line2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.7), Inches(2.95), Inches(2.4), Inches(0.08))
line2.fill.solid(); line2.fill.fore_color.rgb = DARK; line2.line.color.rgb = DARK

b = slide.shapes.add_textbox(Inches(0.9), Inches(5.15), Inches(11.4), Inches(1.3))
text_frame = b.text_frame
for idx, item in enumerate([
    'Atacante pasivo: escucha el canal.',
    'Atacante activo: modifica, inyecta, elimina o repite mensajes.',
    'MITM: Alice y Bob creen que hablan entre sí, pero Eve controla la comunicación.',
]):
    p = text_frame.paragraphs[0] if idx==0 else text_frame.add_paragraph()
    p.text = item
    p.bullet = True
    p.alignment = PP_ALIGN.LEFT
    p.runs[0].font.name = 'Calibri'; p.runs[0].font.size = Pt(18); p.runs[0].font.color.rgb = TEXT

# Slide 4: security goals + primitives
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg_fill(slide)
add_header(slide, 'Objetivos y primitivas')
add_title(slide, 'Seguridad de la información', 'Los bloques básicos de la criptografía')

add_two_columns(
    slide,
    [
        'Confidencialidad: ¿quién puede leer el mensaje?',
        'Integridad: ¿ha sido modificado?',
        'Autenticación: ¿con quién estoy hablando?',
        'No repudio: ¿puede negarse una acción?'
    ],
    [
        'Cifrado: protege el contenido con clave.',
        'Hash: genera una huella digital del contenido.',
        'MAC: verifica integridad y autenticidad compartiendo clave.',
        'Firma digital: autentica el origen y la integridad del mensaje.'
    ],
    title_left='Objetivos',
    title_right='Primitivas'
)

# Slide 5: symmetric/asymmetric and hash
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg_fill(slide)
add_header(slide, 'Criptografía simétrica y asimétrica')
add_title(slide, 'Diferencias clave', 'Qué usa cada enfoque')

left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.1), Inches(4.8))
left_tf = left_box.text_frame
p = left_tf.paragraphs[0]
p.text = 'Criptografía simétrica'
p.runs[0].font.name='Calibri'; p.runs[0].font.size=Pt(22); p.runs[0].font.bold=True; p.runs[0].font.color.rgb=BLUE
for item in [
    'Usa la misma clave para cifrar y descifrar.',
    'Muy rápida y eficiente para grandes volúmenes de datos.',
    'Ejemplo clásico: AES.',
    'Riesgo: la gestión y distribución de claves es crítica.'
]:
    p = left_tf.add_paragraph(); p.text = item; p.bullet = True; p.runs[0].font.name='Calibri'; p.runs[0].font.size=Pt(18); p.runs[0].font.color.rgb = TEXT

right_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.5), Inches(4.8))
right_tf = right_box.text_frame
p = right_tf.paragraphs[0]
p.text = 'Criptografía asimétrica'
p.runs[0].font.name='Calibri'; p.runs[0].font.size=Pt(22); p.runs[0].font.bold=True; p.runs[0].font.color.rgb=BLUE
for item in [
    'Usa clave pública y privada.',
    'Permit e intercambio seguro de claves y autenticación.',
    'Ejemplos: RSA, ECC, ElGamal.',
    'Más lenta, pero fundamental para firmas y certificados.'
]:
    p = right_tf.add_paragraph(); p.text = item; p.bullet = True; p.runs[0].font.name='Calibri'; p.runs[0].font.size=Pt(18); p.runs[0].font.color.rgb = TEXT

# Slide 6: RSA/ECC + Diffie-Hellman
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg_fill(slide)
add_header(slide, 'RSA, ECC y Diffie-Hellman')
add_title(slide, 'Criptografía moderna', 'Fundamentos matemáticos y protocolos clave')
add_bullets_small(slide, [
    'RSA: se basa en la dificultad de factorizar números grandes que son producto de dos primos.',
    'ECC: usa curvas elípticas; ofrece mismo nivel de seguridad con claves más pequeñas.',
    'Diffie-Hellman: permite acordar una clave compartida sin enviarla directamente por el canal.',
    'La clave correcta es combinar estos mecanismos con autenticación para evitar MITM.',
    'En la práctica, todo esto sustenta HTTPS, certificados, firmas digitales y comunicaciones seguras.',
], left=0.9, top=1.7, width=11.6, height=4.5)

# Slide 7: examples + questions
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg_fill(slide)
add_header(slide, 'Ejemplos y cierre')
add_title(slide, 'Puntos clave para recordar', 'Criptografía postcuántica y fundamentos actuales')
add_bullets(slide, [
    'Cifrado resuelve confidencialidad; hash resuelve integridad; firma digital resuelve autenticidad e integridad.',
    'Una clave simétrica es rápida; una clave pública permite autenticar y distribuir claves de forma segura.',
    'AES-GCM y ChaCha20-Poly1305 combinan cifrado y autenticación de forma moderna.',
    'RSA y ECC siguen siendo pilares fundamentales; Diffie-Hellman permite acuerdo de claves sin revelarlas.',
    'La gestión de claves, nonces y certificados es tan importante como el algoritmo elegido.',
])

# final footer
t = slide.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.3))
p = t.text_frame.paragraphs[0]; p.text='Cuestiones de reflexión: ¿qué problema resuelve Diffie-Hellman? ¿qué diferencia hay entre hash y cifrado?'; p.alignment = PP_ALIGN.LEFT; p.runs[0].font.name='Calibri'; p.runs[0].font.size=Pt(12); p.runs[0].font.color.rgb = RGBColor(90, 90, 90)

prs.save(OUT_PATH)
print(f'PowerPoint generated: {OUT_PATH}')
